"""
777c8 Career OS — Supporting Documents API Routes
Handles uploading, text extraction, listing, and deleting supporting documents.
"""

import io
from pathlib import Path
from typing import List

from fastapi import APIRouter, Depends, File, Form, UploadFile, HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from loguru import logger
import fitz  # PyMuPDF
from docx import Document

from app.core.config import settings
from app.core.database import get_db
from app.models import SupportingDocument
from app.schemas import SupportingDocumentResponse

router = APIRouter(prefix="/supporting-documents", tags=["Supporting Documents"])

async def extract_text_from_bytes(data: bytes, file_type: str) -> str:
    """Helper to parse text from pdf, docx, or txt files."""
    file_type = file_type.lower().strip(".")
    if file_type == "pdf":
        try:
            doc = fitz.open(stream=io.BytesIO(data), filetype="pdf")
            text = ""
            for page in doc:
                text += page.get_text()
            doc.close()
            return text
        except Exception as e:
            logger.error(f"Error parsing PDF supporting doc: {e}")
            raise ValueError(f"Failed to parse PDF document: {str(e)}")
    elif file_type == "docx":
        try:
            doc = Document(io.BytesIO(data))
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])
        except Exception as e:
            logger.error(f"Error parsing DOCX supporting doc: {e}")
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")
    else:
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return data.decode("latin-1")
            except Exception:
                raise ValueError("Unsupported or invalid text encoding.")

@router.post("/upload", response_model=SupportingDocumentResponse, status_code=201)
async def upload_supporting_document(
    file: UploadFile = File(...),
    name: str = Form(...),
    doc_type: str = Form("other"),  # transcript, certificate, letter, other
    db: AsyncSession = Depends(get_db),
):
    """Upload a supporting document, extract its text content, and store it securely."""
    file_ext = Path(file.filename).suffix.lower().strip(".")
    supported = {"pdf", "docx", "txt", "md"}
    if file_ext not in supported:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format: .{file_ext}. Supported: PDF, DOCX, TXT, MD",
        )

    # Read uploaded bytes
    content = await file.read()

    # Extract text content
    try:
        extracted_text = await extract_text_from_bytes(content, file_ext)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Save to disk
    target_dir = settings.UPLOADS_DIR / "supporting_docs"
    target_dir.mkdir(parents=True, exist_ok=True)
    file_path = target_dir / f"{name.replace(' ', '_')}_{file.filename}"

    from app.services.settings_service import settings_service
    from app.services.crypto_service import crypto_service

    try:
        system_settings = await settings_service.get_settings(db)
        encrypt = system_settings.encrypt_files
        crypto_service.secure_write_bytes(file_path, content, encrypt=encrypt)
    except Exception as e:
        logger.error(f"Failed to save supporting document: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to write file securely: {str(e)}")

    # Store in database
    doc = SupportingDocument(
        name=name,
        doc_type=doc_type,
        content=extracted_text,
        file_path=str(file_path),
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    return doc

@router.get("/", response_model=List[SupportingDocumentResponse])
async def list_supporting_documents(db: AsyncSession = Depends(get_db)):
    """List all supporting documents (metadata and content)."""
    result = await db.execute(select(SupportingDocument).order_by(SupportingDocument.created_at.desc()))
    return result.scalars().all()

@router.delete("/{doc_id}", status_code=204)
async def delete_supporting_document(doc_id: int, db: AsyncSession = Depends(get_db)):
    """Delete a supporting document by ID."""
    doc = await db.get(SupportingDocument, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    # Attempt to delete file from disk
    if doc.file_path:
        try:
            path = Path(doc.file_path)
            if path.exists():
                path.unlink()
        except Exception as e:
            logger.warning(f"Failed to delete supporting document file on disk: {e}")

    await db.delete(doc)
    await db.commit()
    return None
